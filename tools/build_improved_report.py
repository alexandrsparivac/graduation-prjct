from pathlib import Path
from copy import deepcopy

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Raport_LinguaDomains_imbunatatit.docx"
ASSETS = ROOT / ".report_assets"
ASSETS.mkdir(exist_ok=True)

W, H = 1800, 1060
INK = "#172033"
MUTED = "#5F6B7A"
BLUE = "#285A9B"
BLUE_LIGHT = "#EAF2FD"
GREEN = "#3A7D44"
GREEN_LIGHT = "#EAF6EC"
GOLD = "#9A6800"
GOLD_LIGHT = "#FFF3D5"
RED = "#A74343"
RED_LIGHT = "#FBE9E9"
GRAY = "#E4E9F0"
WHITE = "#FFFFFF"


def font(size, bold=False):
    candidates = [
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/Library/Fonts/Arial.ttf",
        "/System/Library/Fonts/Supplemental/Helvetica.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default()


F_TITLE = font(48, True)
F_HEAD = font(34, True)
F_BODY = font(30)
F_SMALL = font(27)
F_TINY = font(22)


def canvas(title, subtitle):
    im = Image.new("RGB", (W, H), WHITE)
    d = ImageDraw.Draw(im)
    d.rectangle((0, 0, W, 88), fill=INK)
    d.text((48, 20), title, fill=WHITE, font=F_TITLE)
    d.text((48, 105), subtitle, fill=MUTED, font=F_SMALL)
    d.line((48, 142, W - 48, 142), fill=GRAY, width=3)
    return im, d


def rounded(d, box, fill=WHITE, outline=BLUE, width=3, radius=18):
    d.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def text_center(d, box, value, f=F_BODY, fill=INK, spacing=4):
    left, top, right, bottom = box
    maxw = max(40, right - left - 22)
    words = value.split()
    lines, line = [], ""
    for word in words:
        candidate = (line + " " + word).strip()
        if d.textbbox((0, 0), candidate, font=f)[2] <= maxw:
            line = candidate
        else:
            if line:
                lines.append(line)
            line = word
    if line:
        lines.append(line)
    heights = [d.textbbox((0, 0), x, font=f)[3] for x in lines]
    total = sum(heights) + max(0, len(lines)-1)*spacing
    y = top + ((bottom-top)-total)/2
    for line, h in zip(lines, heights):
        width = d.textbbox((0, 0), line, font=f)[2]
        d.text((left+(right-left-width)/2, y), line, fill=fill, font=f)
        y += h + spacing


def arrow(d, xy1, xy2, label=None, dashed=False, fill=INK):
    x1, y1 = xy1
    x2, y2 = xy2
    if dashed:
        steps = 18
        for i in range(steps):
            if i % 2 == 0:
                a, b = i / steps, (i + 1) / steps
                d.line((x1+(x2-x1)*a, y1+(y2-y1)*a, x1+(x2-x1)*b, y1+(y2-y1)*b), fill=fill, width=3)
        # arrow tip
        d.polygon([(x2, y2), (x2-12, y2-7), (x2-12, y2+7)], fill=fill)
    else:
        d.line((x1, y1, x2, y2), fill=fill, width=3)
        if x2 >= x1:
            d.polygon([(x2, y2), (x2-13, y2-7), (x2-13, y2+7)], fill=fill)
        else:
            d.polygon([(x2, y2), (x2+13, y2-7), (x2+13, y2+7)], fill=fill)
    if label:
        midx, midy = (x1+x2)/2, (y1+y2)/2
        bbox = d.textbbox((0, 0), label, font=F_TINY)
        tw, th = bbox[2], bbox[3]
        d.rectangle((midx-tw/2-6, midy-th/2-4, midx+tw/2+6, midy+th/2+4), fill=WHITE)
        d.text((midx-tw/2, midy-th/2), label, fill=INK, font=F_TINY)


def sequence_diagram(path):
    im, d = canvas("Diagramă de secvență – generarea unei lecții", "Scenariul cache-miss; implementare: lesson.html, server.js și Supabase")
    actors = [(130, "Cursant"), (425, "Browser"), (760, "Supabase"), (1110, "Express"), (1510, "Groq")]
    for x, name in actors:
        rounded(d, (x-105, 174, x+105, 230), BLUE_LIGHT, BLUE, 3, 12)
        text_center(d, (x-100, 180, x+100, 225), name, F_BODY)
        d.line((x, 232, x, 950), fill="#94A3B8", width=2)
    # actor stick figure
    d.ellipse((112, 245, 148, 281), outline=INK, width=3)
    d.line((130, 281, 130, 330), fill=INK, width=3)
    d.line((102, 300, 158, 300), fill=INK, width=3)
    d.line((130, 330, 105, 360), fill=INK, width=3)
    d.line((130, 330, 155, 360), fill=INK, width=3)
    arrows = [
        (130, 385, 425, "1  Deschide lecția"),
        (425, 455, 760, "2  SELECT lecție (cheie didactică)"),
        (760, 520, 425, "3  cache absent / incomplet"),
        (425, 590, 1110, "4  POST /api/lesson + JWT"),
        (1110, 655, 760, "5  verifică utilizatorul"),
        (760, 710, 1110, "6  utilizator valid"),
    ]
    for x1, y, x2, label in arrows:
        arrow(d, (x1, y), (x2, y), label)
    # parallel frame
    d.rectangle((970, 744, 1660, 910), outline=BLUE, width=3)
    d.rectangle((970, 744, 1060, 780), fill=BLUE, outline=BLUE)
    d.text((988, 751), "par", fill=WHITE, font=F_SMALL)
    rows = [(780, "7a core"), (820, "7b story"), (860, "7c practice"), (900, "7d quiz")]
    for y, label in rows:
        arrow(d, (1110, y), (1510, y), label)
        arrow(d, (1510, y+14), (1110, y+14), dashed=True, fill=GREEN)
    arrow(d, (1110, 940), (425, 940), "8  merge + validate + răspuns")
    d.text((50, 982), "Notă: persistența în lessons este realizată ulterior de client prin Supabase, după validarea finală.", fill=MUTED, font=F_SMALL)
    im.save(path, quality=95)


def class_diagram(path):
    im, d = canvas("Diagramă de clasă – nucleul de date didactic", "Entități persistente și relații relevante pentru progres și repetarea spațiată")
    boxes = {
        "Profile": (70, 195, 420, 430, ["id : uuid", "full_name : text", "native_language : text", "onboarding_done : bool"]),
        "UserLanguage": (520, 195, 880, 430, ["user_id : uuid", "language_code : text", "level : CEFR"]),
        "UserDomain": (1010, 195, 1370, 430, ["user_id : uuid", "language_code : text", "domain_slug : text", "topics : text[]"]),
        "Lesson": (1400, 195, 1730, 510, ["id : uuid", "language_code : text", "domain_slug : text", "topic : text", "level : CEFR", "content : jsonb"]),
        "UserProgress": (1010, 645, 1370, 900, ["user_id : uuid", "lesson_id : uuid", "completed : bool", "score / total : int"]),
        "UserVocabulary": (380, 645, 835, 980, ["user_id : uuid", "language_code : text", "term : text", "ease : real", "interval_days : int", "due_on : date", "reps / lapses : int"]),
    }
    for name, (l, t, r, b, fields) in boxes.items():
        rounded(d, (l, t, r, b), WHITE, BLUE, 3, 12)
        d.rectangle((l+2, t+2, r-2, t+52), fill=BLUE)
        text_center(d, (l+8, t+4, r-8, t+50), f"«table» {name}", F_BODY, WHITE)
        y=t+67
        for field in fields:
            d.text((l+20, y), field, fill=INK, font=F_SMALL)
            y += 36
    # Associations, deliberately kept outside boxes
    arrow(d, (420, 310), (520, 310), "1     *")
    arrow(d, (880, 310), (1010, 310), "1     *")
    arrow(d, (1370, 350), (1400, 350), dashed=True, fill=GOLD)
    arrow(d, (1560, 510), (1190, 645), "1             *")
    arrow(d, (700, 430), (650, 645), "1     *")
    d.text((70, 1015), "Cheia cache a lecției din schema curentă: (language_code, domain_slug, topic, level).", fill=MUTED, font=F_TINY)
    im.save(path, quality=95)


def component_diagram(path):
    im, d = canvas("Diagramă de componente – granițe de securitate", "Separarea clientului, serverului, serviciilor Supabase și inferenței LLM")
    # zones
    d.rounded_rectangle((55, 175, 560, 905), radius=22, fill="#F8FBFF", outline="#8BA7C9", width=3)
    d.rounded_rectangle((625, 175, 1165, 905), radius=22, fill="#F8FBFF", outline="#8BA7C9", width=3)
    d.rounded_rectangle((1230, 175, 1745, 905), radius=22, fill="#F8FBFF", outline="#8BA7C9", width=3)
    d.text((78, 193), "CLIENT (browser)", fill=BLUE, font=F_HEAD)
    d.text((650, 193), "SERVER APLICATIV", fill=BLUE, font=F_HEAD)
    d.text((1255, 193), "SERVICII CLOUD", fill=BLUE, font=F_HEAD)
    items = [
        ((105, 280, 510, 405), "«component» lesson.html", "validare client, fetch, UI"),
        ((105, 540, 510, 665), "«component» Supabase JS", "Auth, REST, RPC; anon key"),
        ((675, 280, 1115, 435), "«component» Express", "POST /api/lesson\nverifySupabaseUser\ninflight + Promise.all"),
        ((675, 570, 1115, 695), "«component» lesson-prompt.js", "prompt, parse, validate, merge"),
        ((1280, 280, 1690, 440), "«component» Supabase", "Auth + Postgres + RLS\nlessons, progress, vocabulary"),
        ((1280, 570, 1690, 695), "«component» Groq API", "chat completions\nJSON object"),
    ]
    for box, head, body in items:
        rounded(d, box, WHITE, BLUE, 3, 14)
        l,t,r,b=box
        d.text((l+18,t+18),head,fill=INK,font=F_BODY)
        yy=t+62
        for line in body.split("\n"):
            d.text((l+20,yy),line,fill=MUTED,font=F_SMALL); yy+=28
    arrow(d, (510, 340), (675, 340), "HTTPS + JWT")
    arrow(d, (510, 600), (1280, 360), "REST/RPC, RLS", dashed=True, fill=GREEN)
    arrow(d, (1115, 355), (1280, 355), "verificare JWT", dashed=True, fill=GREEN)
    arrow(d, (1115, 630), (1280, 630), "HTTPS + Bearer", fill=RED)
    d.rounded_rectangle((705, 745, 1085, 850), radius=12, fill=RED_LIGHT, outline=RED, width=3)
    text_center(d, (720, 756, 1070, 840), "GROQ_API_KEY rămâne exclusiv în variabilele de mediu ale serverului", F_TINY, RED)
    d.text((72, 948), "RLS aplică izolarea datelor utilizatorilor; cheia anonimă Supabase poate fi expusă clientului prin proiectare.", fill=MUTED, font=F_SMALL)
    im.save(path, quality=95)


def shade_cell(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}'); tcMar.append(node)
        node.set(qn('w:w'), str(v)); node.set(qn('w:type'), 'dxa')


def style_para(p, size=12, bold=False, center=False, italic=False, space_after=6):
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    for r in p.runs:
        r.font.name = 'Times New Roman'
        r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = RGBColor(0, 0, 0)


def add_para(doc, text, size=12, bold=False, center=False, italic=False, before=0, after=6):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_before = Pt(before)
    style_para(p, size, bold, center, italic, after)
    return p


def main():
    seq = ASSETS / "uml_sequence_generation.png"
    classes = ASSETS / "uml_class_core.png"
    comp = ASSETS / "uml_component_security.png"
    sequence_diagram(seq)
    class_diagram(classes)
    component_diagram(comp)

    doc = Document(ROOT / "Raport_LinguaDomains.docx")
    # Corrections grounded in the repository's current implementation.
    corrections = {
        65: "Metodologia de cercetare combină analiza bibliografică, analiza comparativă a platformelor Duolingo, Babbel, Busuu și Memrise, modelarea vizuală UML, prototiparea iterativă și testarea automată (Node test runner, 53 de teste). Baza tehnologică este formată din Node.js 18+, Express 4, Supabase, Groq (lanț configurabil de modele; implicit GPT-OSS și Qwen), JavaScript ES modules și Web Speech API.",
        83: "Groq API, compatibil cu interfața OpenAI chat completions, furnizează inferența LLM. Configurația implicită din server.js utilizează lanțul openai/gpt-oss-120b, openai/gpt-oss-20b și qwen/qwen3.8-27b; lista poate fi înlocuită prin GROQ_MODEL. Parametrii utilizați sunt temperature 0.3 pentru stabilitate didactică, reasoning_effort low pentru a rezerva bugetul de tokeni răspunsului și response_format json_object pentru a impune JSON valid. Bugetele per secțiune sunt calibrate, iar lanțul de fallback permite continuarea transparentă pe modelul următor la epuizarea cotei zilnice, fără eșecul întregii lecții. Detaliile ingineriei promptului sunt tratate în secțiunea 3.3.",
        95: "Din analiza comparativă rezultă că diferențiatorul principal al platformei LinguaDomains constă în granularitatea didactică: cheia persistentă de cache este (limbă, domeniu, subiect, nivel CEFR), iar limba maternă este transmisă ca context la generare pentru explicații și traduceri. Contractul JSON garantează completitudinea conținutului. Implicația acestei alegeri este discutată explicit în secțiunea 2.9, deoarece cheia cache curentă nu distinge limba maternă.",
        111: "Figura 2.2.2 (Onboarding în 4 pași). Secvența interoghează în paralel tabelele languages și domains, randează grilele de opțiuni și validează local pragurile de continuare. La finalizare, browserul execută un upsert în user_languages, rescrie selecțiile din user_domains și marchează profiles.onboarding_done=true; fiecare operație este protejată de RLS. La revenirea cu ?edit=1, pașii sunt precompletați din rândurile existente, iar upsert-ul previne duplicarea limbii. Timpul de răspuns este dominat de două interogări REST, fără apel AI.",
        142: "Figura 2.5.1 (Ciclul de viață al lecției). Stările sunt: Inexistentă → ÎnGenerare → Validată → MemoratăCache → ÎnParcurgere → Finalizată (cu scor) → Regenerată (opțional). Evenimentele sunt: prima vizită (cache-miss), validarea celor 4 secțiuni, eșecul validării (buclă locală pe secțiune), atingerea termenului de 120s (tranziție la EroareTranzitorie cu reîncercare client), quota zilnică (tranziție la Indisponibilă cu mesaj final), parcurgerea pașilor și salvarea scorului. Starea MemoratăCache este stabilă și partajată. Regenerarea solicită actualizarea cache-ului; dacă politica RLS curentă nu permite UPDATE asupra lessons, clientul păstrează lecția nouă pentru vizita curentă și raportează avertismentul în consolă. Vocabularul deja planificat este păstrat prin ignoreDuplicates.",
        152: "Figura 2.6.1 (Onboarding în 4 pași). Nodul inițial este deschiderea /onboarding. Activitățile sunt: A1 alegerea limbii (grilă din languages), A2 alegerea nivelului CEFR (A1–C2 cu descrieri), A3 alegerea domeniilor (grilă din domains), A4 bifarea subiectelor per domeniu. Porțile de decizie verifică selecția minimă înainte de activarea butonului Continuă. Activitatea finală execută persistarea secvențială sub RLS și redirectarea către /dashboard. Ramura alternativă Editare reutilizează aceleași activități cu precompletare. Benzile de responsabilitate Cursant / Browser / Supabase arată că validarea pragurilor este locală, iar fiecare scriere este autorizată separat de RLS.",
    }
    for idx, value in corrections.items():
        p = doc.paragraphs[idx]
        p.text = value
        style_para(p)

    created = []
    heading = add_para(doc, "2.9 Verificarea trasabilității UML față de implementarea curentă", size=14, bold=True, before=12, after=8)
    created.append(heading._p)
    intro = add_para(doc, "Diagramele existente descriu structura generală a sistemului. Pentru a menține raportul aliniat cu codul sursă, această secțiune adaugă trei modele UML focalizate pe traseele cu risc tehnic ridicat: generarea secțională a lecției, nucleul de date didactic și granița de securitate dintre browser, server și serviciile cloud. Fiecare model este trasabil la fișierele server.js, public/lesson.html, lib/lesson-prompt.js, public/js/vocab.js, schema.sql și migrarea 005.")
    created.append(intro._p)

    p = add_para(doc, "2.9.1 Diagrama de secvență pentru generarea lecției", size=13, bold=True, before=7, after=5); created.append(p._p)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(); r.add_picture(str(seq), width=Inches(6.35)); created.append(p._p)
    p = add_para(doc, "Figura 2.9.1 – Diagrama de secvență: cache-miss și generarea paralelă a unei lecții", size=10, center=True, italic=True, after=5); created.append(p._p)
    p = add_para(doc, "Diagrama separă explicit verificarea identității de generarea AI. Browserul verifică mai întâi cache-ul lessons; doar dacă nu există o lecție completă trimite cererea protejată la Express. Serverul validează JWT-ul la Supabase și pornește cele patru secțiuni independente în paralel. Răspunsurile JSON sunt validate și reunite pe server; clientul persistă ulterior rezultatul în cache. Această separare reflectă responsabilitățile din cod și explică de ce o secțiune eșuată poate fi reîncercată fără refacerea întregii lecții."); created.append(p._p)

    p = add_para(doc, "2.9.2 Diagrama de clasă pentru progres și vocabular", size=13, bold=True, before=7, after=5); created.append(p._p)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(); r.add_picture(str(classes), width=Inches(6.35)); created.append(p._p)
    p = add_para(doc, "Figura 2.9.2 – Diagrama de clasă: profil, traseu, lecție, progres și vocabular", size=10, center=True, italic=True, after=5); created.append(p._p)
    p = add_para(doc, "Modelul evidențiază separarea dintre lecția partajată și progresul individual. O lecție memorată în cache poate avea mai multe înregistrări UserProgress, iar vocabularul utilizatorului păstrează propriile scadențe SM-2 independent de regenerarea sursei. Diagrama include cheia de cache care există în schema actuală. Astfel, raportul nu atribuie în mod eronat limba maternă constrângerii unice a tabelei lessons."); created.append(p._p)

    p = add_para(doc, "2.9.3 Diagrama de componente pentru securitate", size=13, bold=True, before=7, after=5); created.append(p._p)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(); r.add_picture(str(comp), width=Inches(6.35)); created.append(p._p)
    p = add_para(doc, "Figura 2.9.3 – Diagrama de componente: delimitarea secretelor și a responsabilităților", size=10, center=True, italic=True, after=5); created.append(p._p)
    p = add_para(doc, "Componenta Express este singura care poate apela Groq cu secretul GROQ_API_KEY. Browserul poate comunica direct cu Supabase numai prin cheia anonimă, iar RLS restrânge accesul la rândurile proprii. Această diagramă completează modelul de desfășurare prin reprezentarea graniței de încredere și prin separarea apelului server–Groq de accesul client–Supabase."); created.append(p._p)

    p = add_para(doc, "Observații de audit și recomandări", size=13, bold=True, before=7, after=5); created.append(p._p)
    p = add_para(doc, "Analiza implementării a identificat două aspecte care trebuie formulate precis în documentație. În primul rând, nativeLanguage intră în prompt și în cheia temporară inflight de pe server, însă nu intră în constrângerea unică sau în SELECT-ul persistent al tabelei lessons. Pentru a garanta explicații în limba maternă corectă pentru fiecare cursant, o versiune viitoare trebuie fie să adauge native_language în cheia cache, fie să separe conținutul studiat de traducerile dependente de utilizator. În al doilea rând, schema curentă definește SELECT și INSERT pentru lessons, nu și UPDATE; prin urmare, regenerarea unei lecții create de alt utilizator poate rămâne doar în memoria sesiunii curente. O politică UPDATE controlată sau un endpoint server-side ar transforma această operație într-una persistentă și auditabilă."); created.append(p._p)

    table = doc.add_table(rows=1, cols=3)
    created.append(table._tbl)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    headers = ["Model UML", "Artefacte trasabile", "Verificare în proiect"]
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value; shade_cell(cell, "1F4E78"); set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.name='Times New Roman'; run.font.size=Pt(10); run.font.bold=True; run.font.color.rgb=RGBColor(255,255,255)
    rows = [
        ("Secvență", "lesson.html; server.js; lib/lesson-prompt.js", "cache, JWT, Promise.all, validateLesson"),
        ("Clasă", "schema.sql; 003_level_minimum.sql; 005_vocabulary.sql", "chei, relații, progres individual, SM-2"),
        ("Componente", "server.js; public/js/supabase.js; public/js/vocab.js", "granița secretului Groq și RLS"),
    ]
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text=value; set_cell_margins(cell); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i % 2 == 0: shade_cell(cell, "F5F9FD")
            for run in cell.paragraphs[0].runs:
                run.font.name='Times New Roman'; run.font.size=Pt(10); run.font.color.rgb=RGBColor(0,0,0)
    p = add_para(doc, "Tabelul 2.9.1 – Matrice de trasabilitate între UML și implementare", size=10, center=True, italic=True, after=8); created.append(p._p)

    # Place the newly built elements before Chapter 3 rather than at document end.
    target = next(
        p._p for p in doc.paragraphs
        if p.text == "3 REALIZAREA APLICAȚIEI UTILIZÂND TEHNOLOGIILE WEB"
    )
    for element in created:
        target.addprevious(element)

    # Update the manually maintained table of contents conservatively.
    toc_insert_after = doc.paragraphs[35]._p
    toc = OxmlElement('w:p')
    r = OxmlElement('w:r'); t = OxmlElement('w:t'); t.text = "2.9 Verificarea trasabilității UML față de implementarea curentă ..... 32"; r.append(t); toc.append(r)
    toc_insert_after.addnext(toc)
    toc_p = next(p for p in doc.paragraphs if p._p is toc)
    style_para(toc_p)

    doc.core_properties.title = "Raport LinguaDomains revizuit"
    doc.core_properties.subject = "Platformă Web pentru învățarea limbilor pe domenii profesionale"
    doc.core_properties.comments = "Raport revizuit cu diagrame UML complementare și verificare de trasabilitate."
    doc.save(OUT)
    print(OUT)


if __name__ == '__main__':
    main()
