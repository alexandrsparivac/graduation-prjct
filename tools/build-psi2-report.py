"""Raport PSI Lab 2 — format UTM, text curgător (fără enumerări seci)."""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

FONT = 'Times New Roman'

doc = Document()
normal = doc.styles['Normal']
normal.font.name = FONT
normal.font.size = Pt(12)
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

for section in doc.sections:
    section.top_margin = Cm(1.9)
    section.bottom_margin = Cm(1.9)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)


def title_line(text, size):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = True
    r.font.name = FONT
    r.font.size = Pt(size)


def h1(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    r = p.add_run(text)
    r.bold = True
    r.font.name = FONT
    r.font.size = Pt(12)


def h2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    r = p.add_run(text)
    r.bold = True
    r.font.name = FONT
    r.font.size = Pt(12)


def body(text):
    p = doc.add_paragraph(text)
    for r in p.runs:
        r.font.name = FONT
        r.font.size = Pt(12)


def lead(title, text):
    p = doc.add_paragraph()
    r = p.add_run(title + ' — ')
    r.bold = True
    r.font.name = FONT
    r.font.size = Pt(12)
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(12)


def grid_table(headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ''
        r = cell.paragraphs[0].add_run(h)
        r.bold = True
        r.font.name = FONT
        r.font.size = Pt(12)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ''
            r = cells[i].paragraphs[0].add_run(val)
            r.font.name = FONT
            r.font.size = Pt(12)
    doc.add_paragraph()


# ---------------- Pagina de titlu ----------------
title_line('MINISTERUL EDUCAȚIEI ȘI CERCETĂRII AL REPUBLICII MOLDOVA', 14)
title_line('Universitatea Tehnică a Moldovei', 13)
title_line('Facultatea Calculatoare, Informatică și Microelectronică', 13)
title_line('Departamentul Ingineria Software', 13)
title_line('RAPORT', 19)
title_line('Lucrarea de laborator nr. 2', 13)
title_line('Disciplina: Proiectarea Sistemelor Informaționale', 13)
title_line('Tema: Dezvoltarea unei platforme pentru învățarea personalizată a limbilor străine în cont', 13)
title_line('Chișinău 2026', 12)
doc.add_page_break()

# ---------------- 1 ----------------
h1('1 DESCRIEREA APLICAȚIEI ȘI A UTILIZATORILOR')
body(
    'LinguaDomains este o aplicație web pentru oricine vrea să învețe o limbă străină pentru '
    'meseria lui: elev, student, angajat sau oricine își schimbă domeniul. Nu este făcută doar '
    'pentru copii și nu are roluri de familie — oricine își face cont învață în același fel. '
    'Utilizatorul își alege limba (una din 25), domeniul în care lucrează (IT, medical, business '
    'și altele), un subiect concret și nivelul, de la începător până la avansat. Din acel moment, '
    'aplicația îi construiește lecții croite exact pe tema lui, îi corectează testele, îi păstrează '
    'cuvintele pentru recapitulare și îi arată progresul într-un panou personal pe care și-l aranjează singur.'
)
h2('1.1 Rolurile: user și admin')
body(
    'În aplicație există exact două roluri. Userul este oricine învață: își face cont, își alege '
    'limbile, domeniile și nivelul, face lecții și teste, repetă vocabularul, își vede progresul, își '
    'mută și ascunde cardurile, își descarcă datele sau își șterge definitiv contul. Vede doar datele '
    'lui și nu se atinge de ale altora. Administratorul este un user marcat în plus, care are o pagină '
    'separată de administrare: caută utilizatori după nume sau email, vede cine e blocat sau cine n-a '
    'terminat pașii de început, blochează conturile problemă scriind motivul și le deblochează la nevoie. '
    'Cine e blocat vede la conectare mesajul de suspendare. Administratorul nu vede parolele nimănui și nu '
    'poate modifica scorurile sau progresul cuiva.'
)
grid_table(
    ['Acțiune', 'User', 'Admin'],
    [
        ('Cont, conectare, recuperare parolă', 'Da', 'Da'),
        ('Lecții, teste, recapitulări', 'Da', 'Da (ca user)'),
        ('Progres, vocabular, carduri proprii', 'Da', 'Da (proprii)'),
        ('Schimbare limbi, teme, setări', 'Da', 'Da'),
        ('Descărcare / ștergere date proprii', 'Da', 'Da'),
        ('Căutare utilizatori', 'Nu', 'Da'),
        ('Blocare / deblocare conturi', 'Nu', 'Da'),
        ('Văzut parole sau modificat scoruri', 'Nu', 'Nu'),
    ],
)

# ---------------- 2 ----------------
h1('2 INTERFAȚA APLICAȚIEI (PROTOTIP FIGMA)')
body(
    'Prototipul Figma se construiește ca design funcțional: frame-uri desktop de 1440px plus variantele '
    'mobile de 390px, cu componente reutilizabile (buton primar și secundar, card, capsulă de header, câmp '
    'de autentificare, switch). Mai jos este fiecare ecran, cu explicația lui; în dreptul fiecăruia se '
    'lipește screenshotul din aplicație.'
)
lead('Autentificarea', 'este ecranul de intrare, cu tot conținutul centrat în pagină. Aici se află câmpurile '
     'pentru email și parolă, butonul de conectare, conectarea cu contul Google și linkul de recuperare a '
     'parolei pentru cei care au uitat-o. Ecranul de înregistrare păstrează același stil, dar câmpurile '
     '(nume, email, parolă, confirmare) stau două câte două pe rând, ca formularul să rămână compact. '
     '[SCREENSHOT: login]')
lead('Alegerea studiului', 'sunt pașii de început, așezați simplu pe verticală: mai întâi limba din cele 25 '
     'disponibile, apoi nivelul, apoi domeniile care te interesează. Totul se salvează dintr-o dată, ca să nu '
     'rămâi cu profilul pe jumătate făcut, iar la final ești dus în pagina de informații. Alte limbi se '
     'adaugă oricând din butonul plus de lângă logo. [SCREENSHOT: onboarding]')
lead('Panoul principal', 'este ecranul de zi cu zi și inima aplicației. Aici stau 15 carduri: seria de zile '
     'consecutive, lecțiile terminate, media scorurilor, cuvintele învățate, obiectivul săptămânii cu butoane '
     'de minus și plus, punctele slabe, calendarul de activitate pe 16 săptămâni, ultimele rezultate și scara '
     'de nivel cu bara de promovare. Pe ecran lat cardurile stau câte patru pe rând, cele late (graficele) '
     'ocupând mai mult loc. Cardurile se trag cu mouse-ul exact unde vrei, iar la hover apar săgeți de mutat '
     'și buton de ascuns; din Setări le aduci înapoi sau le realiniezi cu un singur buton. Ordinea și '
     'vizibilitatea se păstrează în cont, deci le regăsești la fel pe orice dispozitiv. [SCREENSHOT: dashboard]')
lead('Lecțiile', 'încep cu o căsuță de căutare, sub care domeniile stau ca niște sertare: apeși pe numele '
     'domeniului și se deschid cardurile cu subiecte, fiecare cu iconița lui; apeși din nou și se închid, iar '
     'alegerea rămâne memorată și după refresh. Click pe un subiect înseamnă o lecție nouă. '
     '[SCREENSHOT: lecții]')
lead('Lecția în curs', 'curge în pași, unul câte unul pe ecran: mai întâi cele 15 cuvinte cu traduceri, apoi '
     'dialogul din viața reală, textul de citit cu întrebări, gramatica cu exemple, exercițiile de ascultat și '
     'cele de vorbit la microfon, unde primești și o notă pentru pronunție. Testul final are 8 întrebări grilă, '
     'iar ecranul de rezultat îți arată scorul mare, în mijloc; la 8 din 8 ai parte de confetti. Poți naviga cu '
     'săgeata sau cu Enter și poți răspunde cu tastele 1–4, fără mouse. [SCREENSHOT: lecție + rezultat]')
lead('Recapitularea', 'este locul unde ajung singure cuvintele pe care ești pe cale să le uiți. Vezi câte sunt '
     'restante azi, răspunzi cu note, iar aplicația stabilește pentru fiecare când îl revezi, după un algoritm '
     'de repetiție rară. [SCREENSHOT: recapitulare]')
lead('Vocabularul', 'este dulapul cu toate cuvintele salvate din lecții: are căutare, filtre (toate, cuvinte, '
     'expresii, restante) și insigne pentru cele restante. Pe fiecare card vezi traducerea, un exemplu și sursa '
     'exactă — domeniul și subiectul lecției de unde vine cuvântul. [SCREENSHOT: vocabular]')
lead('Setările', 'sunt organizate pe file cu nume clare. La Aspect îți alegi tema (sistem, deschisă sau închisă); '
     'la Limbă îți alegi limba meniurilor din 25; la Limbi de studiu adaugi sau ștergi limbi; la Lecții bifezi cu '
     'switch ce secțiuni să conțină lecțiile; la Widgeturi pornești sau oprești cardurile din panou cu switch; '
     'la Date îți descarci totul ca fișier sau îți ștergi definitiv contul, cu confirmare prin email. '
     '[SCREENSHOT: setări]')
lead('Ghidul și ajutorul', 'sunt pagina de informații: pașii de utilizare și, în 5 pași numerotați, povestea '
     'completă a felului în care se nasc lecțiile. Întrebările frecvente se deschid la click ca un acordeon, '
     'iar Termenii și Confidențialitatea sunt secțiuni simple de citit. [SCREENSHOT: ghid]')
lead('Administrarea', 'o pagină vizibilă doar pentru administratori: câmp de căutare după utilizatori, lista cu '
     'insigne de rol și stare (admin, blocat, pași neterminați), iar la click pe un om se deschid detaliile cu '
     'butoane de blocare (cu motivul scris) sau deblocare. [SCREENSHOT: admin]')

# ---------------- 3 ----------------
h1('3 CUM SE NAȘTE O LECȚIE')
body(
    'Totul pornește când apeși pe un subiect: aplicația cere o lecție nouă pentru combinația ta exactă, adică '
    'limbă plus domeniu plus subiect plus nivel. Înainte de orice, se uită ce ai mai făcut tu, ca lecția să nu '
    'repete cuvintele și subiectele pe care le știi deja, și îi lipește un cod unic. Codul acesta are un rol '
    'practic: dacă apeși de două ori din greșeală sau ai două taburi deschise, a doua cerere o așteaptă pur și '
    'simplu pe prima, în loc să consume resurse de două ori.'
)
body(
    'Scrierea propriu-zisă nu se face dintr-o dată, ci pe bucăți independente, lucrate în paralel: vocabularul '
    'de 15 cuvinte, cele 8 expresii, dialogul de 12–14 replici, textul de citit de 12–16 propoziții cu 3 '
    'întrebări, gramatica cu 4 exemple și 4 exerciții, cele 5 exerciții de ascultat, cele 6 de vorbit și testul '
    'final cu exact 8 întrebări, fiecare cu câte 4 variante. Fiecare bucată e controlată automat de reguli '
    'stricte — număr de elemente, variante complete, cheie de răspuns validă — și ce pică testul se rescrie pe '
    'loc, doar bucata aceea, până iese bine. La final, un control de corectitudine citește toată lecția '
    'cap-coadă ca un corector, iar secțiunile slabe se rescriu integral, nu se peticesc.'
)
body(
    'Momentul-cheie al întregului proces este salvarea în cache. Lecția bună se scrie o singură dată în '
    'memoria comună a platformei, împreună cu o dovadă semnată pe server, care garantează că lecția e '
    'autentică și proaspătă. Din clipa aceea, oricine deschide aceeași combinație primește lecția instant, '
    'fără să mai aștepte generarea. Când deschizi o lecție salvată, în fundal pornește silențios pregătirea '
    'unei variații noi pentru data viitoare, fără să te deranjeze; iar dacă apeși regenerarea, variația din '
    'fundal e oprită mai întâi, ca cele două să nu se calce pe picioare. La finalul quizului, cele 8 răspunsuri '
    'pleacă la server, care le compară cu cheia de răspuns stocată și scrie scorul o singură dată — fiecare '
    'încercare are codul ei unic, ca să nu se poată salva de două ori. Tot atunci cuvintele lecției ajung în '
    'vocabularul tău, iar după 10 teste la rând cu media peste 90%, nivelul tău crește singur cu o treaptă.'
)
body(
    'Un lucru trebuie spus onest: prima deschidere a unui subiect nou durează, pentru că inteligența '
    'artificială folosită este un serviciu gratuit cu cotă limitată, aglomerat la ore de vârf — uneori câteva '
    'minute. În schimb, a doua oară același subiect se deschide instant, din cache, iar tot ce ai deja salvat '
    '(lecții vechi, progres, cuvinte) se deschide mereu instant, indiferent de aglomerație.'
)

# ---------------- 4 ----------------
h1('4 BAZA DE DATE')
body(
    'Toată aplicația vorbește cu o singură bază de date comună, iar regula de bază e simplă: fiecare om vede '
    'și schimbă doar datele lui, în timp ce lecțiile în sine sunt comune tuturor, ca o bibliotecă. Baza e '
    'împărțită în opt sertare. Profilurile țin cine ești — numele, poza, limba meniurilor și dacă ai terminat '
    'pașii de început — și se creează singure în clipa înregistrării. Cataloagele comune spun ce se poate '
    'studia: cele 25 de limbi și domeniile cu subiectele lor, la fel pentru toată lumea. Alegerile tale țin '
    'limbile tale de studiu cu nivelul fiecăreia și domeniile bifate per limbă. Biblioteca de lecții ține '
    'lecțiile validate, una per combinație de limbă, domeniu, subiect și nivel. Traducerile lecțiilor păstrează '
    'aceeași lecție explicată în altă limbă de bază, ca s-o înțeleagă fiecare. Progresul tău ține, pentru '
    'fiecare lecție, unde ai rămas (ca să continui și de pe alt dispozitiv), scorul, data terminării și '
    'varianta ta personală. Testele date țin fiecare quiz corectat, cu scor, limbă și nivel — din ele ies '
    'mediile și promovarea. Vocabularul tău ține fiecare cuvânt cu traducerea, lecția de unde vine și data '
    'următoarei recapitulări.'
)
body(
    'Două automatisme țin totul coerent fără ca aplicația să le ceară de fiecare dată: la terminarea unui '
    'quiz, testul se trece singur în istoric, iar la 10 teste la rând cu media peste 90%, nivelul crește singur.'
)
h2('4.1 Ce face fiecare pagină cu baza de date')
body(
    'Autentificarea citește contul și profilul, iar la înregistrare le creează pe amândouă; la recuperarea '
    'parolei emite linkul cu timp limitat. Alegerea studiului citește cataloagele de limbi și domenii, apoi '
    'scrie dintr-o dată limba, nivelul și domeniile — totul sau nimic, ca să nu rămâi cu profilul pe jumătate '
    'făcut. Panoul principal doar citește: alegerile tale, lecțiile din bibliotecă, progresul cu scoruri, '
    'testele date și vocabularul, pe care le combină în cele 15 carduri; în schimb, ordinea cardurilor, '
    'vizibilitatea lor și obiectivul săptămânii se salvează în cont, nu în browser. Pagina de lecții citește '
    'domeniile tale și caută în bibliotecă ce lecții există deja pentru limba și nivelul tău. Lecția în curs '
    'citește lecția (din bibliotecă sau proaspăt generată) și scrie continuu unde ai rămas, ca să poți continua '
    'și de pe alt dispozitiv; la final scrie scorul prin funcția de corectare, salvează cuvintele noi și golește '
    'ciorna, iar lecția proaspătă ajunge mai întâi în bibliotecă, cu dovada semnată că e autentică. '
    'Recapitularea citește cuvintele restante de azi și scrie înapoi noua dată de repetare după fiecare notă '
    'pe care o dai. Vocabularul citește toate cuvintele tale din limba activă, iar căutarea și filtrele se fac '
    'pe loc, în pagină, fără alte cereri. Setările scriu limba meniurilor în profil și limbile de studiu, '
    'șterg o limbă cu tot cu lecțiile ei personale (scorurile trecute rămân), exportă totul ca fișier sau '
    'șterg definitiv contul cu tot ce are. Ghidul și ajutorul nu ating deloc baza de date — sunt pagini de '
    'citit. Administrarea caută printre profiluri și scrie blocarea sau deblocarea conturilor, cu motiv.'
)

# ---------------- 5 ----------------
h1('5 CERINȚE FUNCȚIONALE')
grid_table(
    ['Cod', 'Cerința', 'Descriere'],
    [
        ('RF-01', 'Cont cu email și parolă', 'Înregistrare cu nume, email și parolă de minim 8 caractere; conectare cu verificare; datele fiecăruia separate.'),
        ('RF-02', 'Conectare cu Google', 'Intrare cu contul Google, fără parolă nouă; contul se creează singur la prima conectare.'),
        ('RF-03', 'Recuperare parolă', 'Link pe email, cu timp limitat, pentru setarea unei parole noi.'),
        ('RF-04', 'Alegerea studiului', 'Limbă, nivel și domenii salvate dintr-o dată; limbi multiple în paralel, comutare memorată pe dispozitive.'),
        ('RF-05', 'Lecții făcute pe loc', 'Construite pe subiectul cerut, verificate bucată cu bucată, salvate o singură dată.'),
        ('RF-06', 'Test corectat pe server', '8 întrebări comparate cu cheia de răspuns; scor salvat o singură dată.'),
        ('RF-07', 'Cuvinte care nu se uită', 'Salvate automat și readuse la repetat când ești pe cale să le uiți.'),
        ('RF-08', 'Panou cu 15 carduri', 'Mutate, ascunse și realiniate, cu totul salvat în cont.'),
        ('RF-09', 'Nivel care crește singur', '10 teste la rând cu media peste 90% duc la nivelul următor.'),
        ('RF-10', 'Datele tale, ale tale', 'Descărcare completă ca fișier și ștergere definitivă a contului.'),
    ],
)

# ---------------- 6 ----------------
h1('6 CONSTRÂNGERI')
body(
    'Din punct de vedere tehnic, aplicația e un site web: merge direct din browser, fără să instalezi nimic. '
    'Lecțiile sunt făcute de un serviciu gratuit de inteligență artificială, care are cotă limitată, așa că '
    'la ore aglomerate o lecție complet nouă poate dura câteva minute; în schimb, tot ce ai deja salvat — '
    'lecții vechi, progres, cuvinte — se deschide mereu instant. Cheile secrete stau doar pe server și nu '
    'ajung niciodată în browser.'
)
body(
    'Ca sistem de operare merge orice ai: Windows, macOS, Linux pe calculator sau Android și iOS pe telefon '
    'și tabletă, atâta timp cât ai un browser modern la zi — Chrome, Edge, Safari sau Firefox — cu JavaScript '
    'activ. Pentru exercițiile vorbite îți trebuie microfon și sunet pornit. Nu-ți trebuie calculator '
    'puternic și nici mult spațiu: o filă obișnuită de browser e de ajuns, iar o lecție salvată ocupă mai '
    'puțin decât o poză.'
)

# ---------------- Bibliografie ----------------
h1('BIBLIOGRAFIE')
for ref in [
    'Supabase. „Documentation". https://supabase.com/docs',
    'Groq. „Documentation". https://console.groq.com/docs',
    'React. „Documentation". https://react.dev',
    'Vite. „Guide". https://vite.dev/guide',
    'MDN Web Docs. https://developer.mozilla.org',
]:
    p = doc.add_paragraph()
    r = p.add_run(ref)
    r.font.name = FONT
    r.font.size = Pt(12)

# ---------------- Semnături ----------------
doc.add_paragraph()
table = doc.add_table(rows=3, cols=2)
table.style = 'Table Grid'
table.alignment = WD_TABLE_ALIGNMENT.CENTER
for row, vals in [
    (0, ('A efectuat:', 'st. gr. TI-234 Sparivac Alexandr')),
    (1, ('A verificat:', 'asist. univ. Cojocaru Svetlana')),
    (2, ('A verificat coordonator:', 'asist. univ. Astafi Valentina')),
]:
    for i, val in enumerate(vals):
        cell = table.rows[row].cells[i]
        cell.text = ''
        r = cell.paragraphs[0].add_run(val)
        r.font.name = FONT
        r.font.size = Pt(12)

doc.save('Raport_PSI_Lab2_LinguaDomains.docx')
print('saved Raport_PSI_Lab2_LinguaDomains.docx')
